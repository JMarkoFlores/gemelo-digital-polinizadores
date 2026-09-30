import os
import re
import time
from pathlib import Path
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# Mapping of file paths to technical descriptions
FILE_DESCRIPTIONS = {
    ".env.example": "Plantilla con variables de entorno para configuración local y Docker (puertos, conexión PostgreSQL, secretos JWT).",
    ".gitignore": "Reglas de exclusión de Git para artefactos de compilación, cachés de Python, node_modules y entornos virtuales.",
    "docker-compose.yml": "Orquestación multi-contenedor de servicios: Backend FastAPI, Frontend React/Vite, PostgreSQL y Streamlit.",
    "package.json": "Manifiesto raíz con scripts de gestión, arranque simultáneo y dependencias globales del proyecto.",
    "metadata.json": "Metadatos informativos y configuración descriptiva del proyecto de gemelos digitales agroecológicos.",
    "README.md": "Documentación técnica principal, arquitectura general, requisitos previos e instrucciones de despliegue.",
    "CAMBIOS_FUTUROS.md": "Hoja de ruta, backlog de funcionalidades pendientes y optimizaciones arquitectónicas planificadas.",
    "patch_app.py": "Script utilitario para comprobación, parcheo de endpoints y sincronización de modelos entre componentes.",
    
    # Backend
    "backend/Dockerfile": "Imagen Docker para despliegue del servicio FastAPI con dependencias Python, compiladores y Uvicorn.",
    "backend/requirements.txt": "Lista de dependencias Python del backend (FastAPI, SQLAlchemy, psycopg2, PyJWT, scikit-learn, numpy).",
    "backend/.dockerignore": "Reglas de exclusión de archivos no necesarios durante el build de la imagen Docker del backend.",
    "backend/app/__init__.py": "Inicializador del paquete Python principal de la aplicación backend FastAPI.",
    "backend/app/config.py": "Carga centralizada de configuraciones y variables de entorno mediante Pydantic BaseSettings.",
    "backend/app/database.py": "Inicialización del motor SQLAlchemy, declarative base y fábrica de sesiones (SessionLocal) para PostgreSQL.",
    "backend/app/dependencies.py": "Inyección de dependencias para FastAPI: conexión de base de datos, validación de token JWT y roles de usuario.",
    "backend/app/main.py": "Punto de entrada de FastAPI: middleware CORS, registro de routers, ciclo de vida de la app y endpoints base.",
    "backend/app/schemas.py": "Esquemas Pydantic para validación y serialización de usuarios, autenticación, simulaciones y optimizaciones.",
    "backend/app/schemas_reports.py": "Esquemas Pydantic específicos para parámetros de consulta, filtrado y respuestas de reportes ejecutivos.",
    "backend/app/security.py": "Funciones criptográficas: hashing seguro de contraseñas con bcrypt y generación/verificación de tokens JWT.",
    "backend/app/seed.py": "Script de inicialización de datos base: usuarios iniciales (admin y cliente) y escenarios de prueba.",
    "backend/app/models/__init__.py": "Exportador centralizado de modelos ORM para integración con SQLAlchemy y migraciones.",
    "backend/app/models/user.py": "Modelo ORM de entidad de usuario: credenciales, roles de acceso (admin/client) y fechas de auditoría.",
    "backend/app/models/simulation.py": "Modelo ORM de entidad de simulación: metadatos de parcelas, configuración de capas y resultados métricos.",
    "backend/app/routers/__init__.py": "Inicializador del paquete de routers de endpoints de FastAPI.",
    "backend/app/routers/admin_reports.py": "Router FastAPI con endpoints para generación, filtrado, exportación y KPIs de reportes administrativos.",
    "backend/app/services/__init__.py": "Inicializador del paquete de lógica de negocio y servicios especializados.",
    "backend/app/services/chat_service.py": "Servicio de asistente conversacional inteligente para consultas agroecológicas y orientación agronómica.",
    "backend/app/services/model_store.py": "Gestor de persistencia, versionado y carga de modelos de Machine Learning entrenados en disco.",
    "backend/app/services/optimization.py": "Algoritmos de optimización espacial para distribución óptima de cultivos y hábitats de polinizadores.",
    "backend/app/services/recommendation_service.py": "Motor experto de recomendaciones agroecológicas basadas en índices de diversidad y condiciones de suelo.",
    "backend/app/services/report_general_service.py": "Cálculo de métricas agregadas, resúmenes estadísticos e indicadores de rendimiento de simulaciones.",
    "backend/app/services/report_general_export_service.py": "Servicio de exportación y preparación de datos consolidados para reportes administrativos.",
    "backend/app/services/terrain_service.py": "Procesamiento de matrices geoespaciales de terreno, cálculo de pendientes, elevación y zonificación.",
    
    # Frontend
    "frontend/Dockerfile": "Imagen Docker multi-etapa para empaquetado de producción con Vite y servicio web con Nginx/Node.",
    "frontend/.dockerignore": "Reglas de exclusión para evitar incluir archivos innecesarios en la construcción de la imagen del frontend.",
    "frontend/package.json": "Manifiesto de dependencias npm: React 18, Vite, Tailwind CSS, Lucide icons, Leaflet, Three.js y scripts.",
    "frontend/vite.config.js": "Configuración del bundler Vite: plugin React, alias de rutas y proxy inverso hacia el backend en desarrollo.",
    "frontend/tailwind.config.js": "Configuración del framework Tailwind CSS: colores de marca, tipografías y extensiones de diseño.",
    "frontend/postcss.config.js": "Configuración de procesamiento PostCSS con plugins de Tailwind CSS y Autoprefixer.",
    "frontend/index.html": "Plantilla HTML principal de la Single Page Application (SPA) con metadatos y contenedor raíz.",
    "frontend/src/main.jsx": "Punto de entrada de React: inicialización del DOM, React StrictMode y proveedores globales de contexto.",
    "frontend/src/App.jsx": "Enrutamiento principal de la aplicación con React Router, rutas protegidas y diseño responsivo.",
    "frontend/src/index.css": "Estilos globales de la aplicación, importaciones de Tailwind y personalizaciones de componentes.",
    "frontend/src/state/AuthContext.jsx": "Contexto global de autenticación: persistencia de JWT, login, logout, roles y permisos de acceso.",
    "frontend/src/state/UiContext.jsx": "Contexto de interfaz de usuario para gestión de notificaciones toast, modales y temas visuales.",
    "frontend/src/lib/api.js": "Cliente de red centralizado con interceptores automáticos de autenticación JWT y manejo de errores.",
    "frontend/src/lib/i18n.js": "Motor y diccionario de internacionalización para soporte bilingüe de textos en la interfaz.",
    "frontend/src/lib/landUseColors.js": "Paleta cromática y nomenclatura estándar para cada categoría de uso y cobertura de suelo.",
    "frontend/src/lib/leaflet.js": "Configuración de iconos, capas base y adaptadores de mapas interactivos con la librería Leaflet.",
    "frontend/src/lib/exporters.js": "Utilidades en el cliente para exportar datos y tablas de simulación a formatos CSV y JSON.",
    "frontend/src/lib/generalReportExporters.js": "Exportador cliente de reportes consolidados con generación de documentos PDF, Excel y Word.",
    
    # Frontend Components
    "frontend/src/components/AppShell.jsx": "Estructura principal de navegación: barra lateral retráctil, barra superior y perfil de usuario.",
    "frontend/src/components/AuthCard.jsx": "Tarjeta estilizada contenedora para los formularios de autenticación y registro.",
    "frontend/src/components/ChatWidget.jsx": "Widget interactivo de chat flotante con historial conversacional y streaming del asistente IA.",
    "frontend/src/components/ComparisonMapsCard.jsx": "Tarjeta de visualización comparativa de mapas lado a lado (escenario actual vs. optimizado).",
    "frontend/src/components/EmptyState.jsx": "Componente visual informativo para vistas sin registros o estados vacíos de simulación.",
    "frontend/src/components/ErrorBoundary.jsx": "Límite de captura de errores en tiempo de ejecución de React para evitar caídas de la aplicación.",
    "frontend/src/components/ExpandableMapCard.jsx": "Tarjeta de mapa interactivo con funcionalidad de pantalla completa, capas y leyendas dinámicas.",
    "frontend/src/components/FormField.jsx": "Campo de formulario reutilizable con etiquetas, control de validación y mensajes de retroalimentación.",
    "frontend/src/components/FullScreenLoader.jsx": "Pantalla de carga global con indicador animado para operaciones prolongadas o inicialización.",
    "frontend/src/components/LandscapeDiorama3D.jsx": "Visualizador 3D interactivo implementado con Three.js que renderiza el relieve y cobertura vegetal en 3D.",
    "frontend/src/components/LandUseDistributionCards.jsx": "Tarjetas estadísticas de desglose porcentual de usos de suelo con barras de progreso comparativas.",
    "frontend/src/components/MapSelectionCard.jsx": "Selector cartográfico interactivo que permite delimitar áreas de interés mediante polígonos.",
    "frontend/src/components/MetricCard.jsx": "Tarjeta de indicador clave (KPI) con valor numérico, variación porcentual e icono descriptivo.",
    "frontend/src/components/PanelCard.jsx": "Contenedor base para paneles de información con sombras suaves y esquinas redondeadas.",
    "frontend/src/components/ProtectedRoute.jsx": "Guardia de enrutamiento que restringe acceso por autenticación y nivel de rol requerido.",
    "frontend/src/components/RecommendationAccordion.jsx": "Acordeón desplegable con recomendaciones agroecológicas organizadas por nivel de prioridad.",
    "frontend/src/components/ReportExportBar.jsx": "Barra de acciones rápidas para descarga y exportación de reportes en múltiples formatos.",
    "frontend/src/components/ResultsDashboard.jsx": "Tablero integrado de resultados de simulación con métricas agroecológicas, gráficos y mapas.",
    "frontend/src/components/ScenarioPanel.jsx": "Panel interactivo para configurar variables agronómicas y parámetros ambientales del escenario.",
    "frontend/src/components/SimulationHistoryList.jsx": "Listado histórico de simulaciones previas con filtros, fecha, usuario y acciones de recarga.",
    "frontend/src/components/SpinnerBlock.jsx": "Indicador visual de espera contextual con animación giratoria y etiqueta de progreso.",
    "frontend/src/components/StatusBanner.jsx": "Banner de notificación para comunicar estados del sistema, alertas y advertencias operacionales.",
    
    # Frontend Pages
    "frontend/src/pages/LoginPage.jsx": "Página de inicio de sesión con validación de credenciales y redirección por rol.",
    "frontend/src/pages/RegisterPage.jsx": "Página de registro de nuevos usuarios en el sistema con verificación de campos.",
    "frontend/src/pages/ClientDashboard.jsx": "Vista principal del panel para clientes con acceso directo a optimizaciones y simulaciones.",
    "frontend/src/pages/ClientOptimizePage.jsx": "Página principal para configuración, ejecución y análisis de optimización de paisajes agroecológicos.",
    "frontend/src/pages/ClientHistoryPage.jsx": "Historial detallado y auditoría de todas las simulaciones ejecutadas por el cliente.",
    "frontend/src/pages/AdminDashboard.jsx": "Vista principal del panel administrativo para supervisión general de la plataforma.",
    "frontend/src/pages/AdminHomePage.jsx": "Dashboard central de métricas globales, actividad de usuarios y estado de servidores.",
    "frontend/src/pages/AdminSimulationsPage.jsx": "Gestión, monitoreo y auditoría exhaustiva de todas las simulaciones de la base de datos.",
    "frontend/src/pages/AdminReportsPage.jsx": "Módulo administrativo avanzado para generación de reportes ejecutivos y comparativas.",
    "frontend/src/pages/AdminUsersPage.jsx": "Administración de usuarios del sistema: altas, edición de roles, permisos y estados.",
    
    # Streamlit App
    "streamlit_app/Dockerfile": "Imagen Docker para el microservicio Streamlit del motor analítico de Gemelos Digitales.",
    "streamlit_app/.dockerignore": "Reglas de exclusión de archivos temporales para la imagen Docker de Streamlit.",
    "streamlit_app/requirements.txt": "Dependencias de ciencia de datos: Streamlit, Scikit-Learn, PyTorch, Plotly, SciPy, Matplotlib.",
    "streamlit_app/.streamlit/config.toml": "Configuración de tema visual, paleta de colores y puertos del servidor de Streamlit.",
    "streamlit_app/app.py": "Aplicación completa del Gemelo Digital interactivo con paneles de simulación y visualización avanzada.",
    "streamlit_app/training.py": "Pipeline base de entrenamiento y evaluación comparativa de algoritmos de Machine Learning.",
    "streamlit_app/advanced_training.py": "Pipeline de entrenamiento avanzado: modelos híbridos, redes neuronales y ensembles predictivos.",
    "streamlit_app/data_pipeline.py": "Pipeline de ingesta, limpieza, normalización e ingeniería de características espaciales del terreno.",
    "streamlit_app/hyperparameter_tuning.py": "Módulo de optimización y ajuste fino de hiperparámetros con GridSearch y validación cruzada.",
    "streamlit_app/pollinator_abm.py": "Modelo Basado en Agentes (ABM) para simular dinámicas de forrajeo y supervivencia de colonias de abejas.",
    "streamlit_app/reports.py": "Generador de reportes técnicos estadísticos y gráficos de desempeño de modelos en Streamlit.",
    "streamlit_app/robust_tests.py": "Batería de pruebas estadísticas robustas (Friedman, Wilcoxon, análisis de varianza y bootstrap).",
    "streamlit_app/test_traceability_pipeline.py": "Pruebas de verificación de trazabilidad extremo a extremo de datos, modelos y artefactos.",
    "streamlit_app/test_tuning_benchmark.py": "Pruebas de benchmarking y comparación de tiempos de ejecución de ajuste de hiperparámetros.",
    
    # Langflow & Scratch
    "langflow/flows/asistente_recomendaciones_agroecologicas.json": "Definición del flujo conversacional de IA en formato JSON para el agente agronómico en Langflow.",
    "scratch/test_nemenyi.py": "Script estadístico para cálculo del test post-hoc de diferencias críticas de Nemenyi.",
}

# Categorize files into 5 logical modules
MODULES = [
    {
        "id": 1,
        "name": "MÓDULO 1: CONFIGURACIÓN RAÍZ E INFRAESTRUCTURA DOCKER",
        "description": "Archivos de configuración global, orquestación de contenedores, variables de entorno y documentación técnica del proyecto.",
        "color": "1E3A8A",
        "files": [
            "docker-compose.yml",
            ".env.example",
            ".gitignore",
            "package.json",
            "metadata.json",
            "README.md",
            "CAMBIOS_FUTUROS.md",
            "patch_app.py",
        ]
    },
    {
        "id": 2,
        "name": "MÓDULO 2: BACKEND - API REST, MODELOS ORM Y SERVICIOS ML (FASTAPI)",
        "description": "Núcleo de servicios backend desarrollado en FastAPI y SQLAlchemy. Incluye autenticación JWT, modelos de datos en PostgreSQL, endpoints de administración y servicios de optimización agroecológica.",
        "color": "0D9488",
        "files": [
            "backend/Dockerfile",
            "backend/requirements.txt",
            "backend/.dockerignore",
            "backend/app/__init__.py",
            "backend/app/config.py",
            "backend/app/database.py",
            "backend/app/dependencies.py",
            "backend/app/security.py",
            "backend/app/seed.py",
            "backend/app/main.py",
            "backend/app/schemas.py",
            "backend/app/schemas_reports.py",
            "backend/app/models/__init__.py",
            "backend/app/models/user.py",
            "backend/app/models/simulation.py",
            "backend/app/routers/__init__.py",
            "backend/app/routers/admin_reports.py",
            "backend/app/services/__init__.py",
            "backend/app/services/model_store.py",
            "backend/app/services/terrain_service.py",
            "backend/app/services/optimization.py",
            "backend/app/services/recommendation_service.py",
            "backend/app/services/report_general_service.py",
            "backend/app/services/report_general_export_service.py",
            "backend/app/services/chat_service.py",
        ]
    },
    {
        "id": 3,
        "name": "MÓDULO 3: FRONTEND - INTERFAZ DE USUARIO, COMPONENTES Y VISUALIZACIÓN 3D (REACT / VITE)",
        "description": "Aplicación cliente moderna desarrollada en React 18, Vite y Tailwind CSS. Cuenta con dioramas 3D interactivos en Three.js, visualización cartográfica en Leaflet, dashboards administrativos y flujos de simulación.",
        "color": "4F46E5",
        "files": [
            "frontend/Dockerfile",
            "frontend/.dockerignore",
            "frontend/package.json",
            "frontend/vite.config.js",
            "frontend/tailwind.config.js",
            "frontend/postcss.config.js",
            "frontend/index.html",
            "frontend/src/main.jsx",
            "frontend/src/App.jsx",
            "frontend/src/index.css",
            "frontend/src/state/AuthContext.jsx",
            "frontend/src/state/UiContext.jsx",
            "frontend/src/lib/api.js",
            "frontend/src/lib/i18n.js",
            "frontend/src/lib/landUseColors.js",
            "frontend/src/lib/leaflet.js",
            "frontend/src/lib/exporters.js",
            "frontend/src/lib/generalReportExporters.js",
            "frontend/src/components/AppShell.jsx",
            "frontend/src/components/AuthCard.jsx",
            "frontend/src/components/ChatWidget.jsx",
            "frontend/src/components/ComparisonMapsCard.jsx",
            "frontend/src/components/EmptyState.jsx",
            "frontend/src/components/ErrorBoundary.jsx",
            "frontend/src/components/ExpandableMapCard.jsx",
            "frontend/src/components/FormField.jsx",
            "frontend/src/components/FullScreenLoader.jsx",
            "frontend/src/components/LandscapeDiorama3D.jsx",
            "frontend/src/components/LandUseDistributionCards.jsx",
            "frontend/src/components/MapSelectionCard.jsx",
            "frontend/src/components/MetricCard.jsx",
            "frontend/src/components/PanelCard.jsx",
            "frontend/src/components/ProtectedRoute.jsx",
            "frontend/src/components/RecommendationAccordion.jsx",
            "frontend/src/components/ReportExportBar.jsx",
            "frontend/src/components/ResultsDashboard.jsx",
            "frontend/src/components/ScenarioPanel.jsx",
            "frontend/src/components/SimulationHistoryList.jsx",
            "frontend/src/components/SpinnerBlock.jsx",
            "frontend/src/components/StatusBanner.jsx",
            "frontend/src/pages/LoginPage.jsx",
            "frontend/src/pages/RegisterPage.jsx",
            "frontend/src/pages/ClientDashboard.jsx",
            "frontend/src/pages/ClientOptimizePage.jsx",
            "frontend/src/pages/ClientHistoryPage.jsx",
            "frontend/src/pages/AdminDashboard.jsx",
            "frontend/src/pages/AdminHomePage.jsx",
            "frontend/src/pages/AdminSimulationsPage.jsx",
            "frontend/src/pages/AdminReportsPage.jsx",
            "frontend/src/pages/AdminUsersPage.jsx",
        ]
    },
    {
        "id": 4,
        "name": "MÓDULO 4: MOTOR DE GEMELOS DIGITALES Y MACHINE LEARNING (STREAMLIT / ABM)",
        "description": "Motor analítico de gemelos digitales y simulaciones biofísicas implementado en Streamlit y Python científico. Incluye modelos basados en agentes (ABM) de polinizadores, pipelines de datos y benchmarking estadístico.",
        "color": "EA580C",
        "files": [
            "streamlit_app/Dockerfile",
            "streamlit_app/.dockerignore",
            "streamlit_app/requirements.txt",
            "streamlit_app/.streamlit/config.toml",
            "streamlit_app/app.py",
            "streamlit_app/training.py",
            "streamlit_app/advanced_training.py",
            "streamlit_app/data_pipeline.py",
            "streamlit_app/hyperparameter_tuning.py",
            "streamlit_app/pollinator_abm.py",
            "streamlit_app/reports.py",
            "streamlit_app/robust_tests.py",
            "streamlit_app/test_traceability_pipeline.py",
            "streamlit_app/test_tuning_benchmark.py",
        ]
    },
    {
        "id": 5,
        "name": "MÓDULO 5: AGENTES DE IA (LANGFLOW) Y SCRIPTS DE SOPORTE",
        "description": "Definiciones de agentes inteligentes y flujos de razonamiento agroecológico en Langflow, además de scripts de soporte estadístico.",
        "color": "7C3AED",
        "files": [
            "langflow/flows/asistente_recomendaciones_agroecologicas.json",
            "scratch/test_nemenyi.py",
        ]
    }
]

def get_language(filename):
    ext = Path(filename).suffix.lower()
    mapping = {
        ".py": "Python",
        ".jsx": "React JSX (JavaScript)",
        ".js": "JavaScript",
        ".css": "CSS (Tailwind)",
        ".html": "HTML5",
        ".json": "JSON",
        ".yml": "YAML",
        ".yaml": "YAML",
        ".md": "Markdown",
        ".toml": "TOML",
        ".txt": "Texto Plano / Requirements",
        ".example": "Variables de Entorno",
        ".dockerignore": "Configuración Docker",
    }
    if Path(filename).name == "Dockerfile":
        return "Dockerfile"
    return mapping.get(ext, ext or "Texto")

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_cell_background(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_table_borders(table, color_hex="E2E8F0"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="{color_hex}"/>
            <w:left w:val="none"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color_hex}"/>
            <w:right w:val="none"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>
            <w:insideV w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def clean_text(text):
    # Remove XML-illegal characters
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', text)

def build_word_document():
    print("Iniciando la construcción del documento Word...")
    t0 = time.time()
    
    doc = docx.Document()
    
    # Configure page setup: Standard margins 1.8 cm (~0.71 inches)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        
        # Configure Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Gemelo Digital Agroecológico — Código Fuente Completo del Proyecto")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.0)
        hrun.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
        
        # Configure Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Documentación Técnica Oficial de Código Fuente — Proyecto Gemelos Digitales")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.0)
        frun.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    # Calculate overall stats
    total_files = 0
    total_lines = 0
    total_bytes = 0
    file_records = []
    
    for mod in MODULES:
        for rel_path in mod["files"]:
            p = Path(rel_path)
            if not p.exists():
                print(f"Advertencia: Archivo no encontrado: {rel_path}")
                continue
            try:
                raw = p.read_text(encoding="utf-8")
            except UnicodeDecodeError:
                raw = p.read_text(encoding="latin-1")
            
            cleaned = clean_text(raw)
            lines = cleaned.splitlines()
            size = p.stat().st_size
            
            total_files += 1
            total_lines += len(lines)
            total_bytes += size
            
            desc = FILE_DESCRIPTIONS.get(rel_path, "Archivo de código fuente del sistema.")
            lang = get_language(rel_path)
            
            file_records.append({
                "module_id": mod["id"],
                "module_name": mod["name"],
                "module_color": mod["color"],
                "rel_path": rel_path,
                "language": lang,
                "lines_count": len(lines),
                "size_bytes": size,
                "description": desc,
                "content": cleaned,
                "lines": lines
            })

    print(f"Estadísticas calculadas: {total_files} archivos, {total_lines:,} líneas de código, {total_bytes / (1024*1024):.2f} MB")

    # ==========================================
    # 1. PORTADA PROFESIONAL
    # ==========================================
    p_pre = doc.add_paragraph()
    p_pre.paragraph_format.space_before = Pt(36)
    p_pre.paragraph_format.space_after = Pt(8)
    r_tag = p_pre.add_run("EXPEDIENTE TÉCNICO DE SOFTWARE")
    r_tag.font.name = "Calibri"
    r_tag.font.size = Pt(11)
    r_tag.font.bold = True
    r_tag.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
    
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(12)
    r_title = p_title.add_run("SISTEMA DE GEMELOS DIGITALES AGROECOLÓGICOS PARA POLINIZADORES")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(26)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(24)
    r_sub = p_sub.add_run("Compilación Integral y Exhaustiva del Código Fuente — Separado por Archivo")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(14)
    r_sub.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    # Decorative separator bar
    p_bar = doc.add_paragraph()
    p_bar.paragraph_format.space_before = Pt(0)
    p_bar.paragraph_format.space_after = Pt(24)
    r_bar = p_bar.add_run("―" * 58)
    r_bar.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
    r_bar.font.bold = True

    # Summary Stats Table on Cover Page
    stat_tbl = doc.add_table(rows=6, cols=2)
    stat_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    stat_tbl.autofit = False

    stat_data = [
        ("Proyecto", "Gemelo Digital Agroecológico de Polinizadores"),
        ("Repositorio", "JMarkoFlores / gemelo-digital-polinizadores"),
        ("Total de Archivos Incluidos", f"{total_files} archivos (100% del código fuente)"),
        ("Volumen de Código", f"{total_lines:,} líneas de código"),
        ("Stack Tecnológico Principal", "Python 3.13 (FastAPI, Streamlit, PyTorch, Scikit-Learn) + React 18 / Vite + Tailwind CSS + Three.js + PostgreSQL + Langflow + Docker"),
        ("Fecha de Compilación", "Septiembre 2026")
    ]

    col_widths = [Inches(2.4), Inches(4.7)]
    for row_idx, (label, val) in enumerate(stat_data):
        row = stat_tbl.rows[row_idx]
        
        # Col 0: Label
        cell_0 = row.cells[0]
        cell_0.width = col_widths[0]
        set_cell_margins(cell_0, top=90, bottom=90, left=120, right=120)
        set_cell_background(cell_0, "F1F5F9")
        p0 = cell_0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.font.name = "Calibri"
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        
        # Col 1: Value
        cell_1 = row.cells[1]
        cell_1.width = col_widths[1]
        set_cell_margins(cell_1, top=90, bottom=90, left=120, right=120)
        set_cell_background(cell_1, "FFFFFF" if row_idx % 2 == 0 else "F8FAFC")
        p1 = cell_1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(val)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    set_table_borders(stat_tbl, "CBD5E1")

    # Architecture Overview Box
    p_arch_title = doc.add_paragraph()
    p_arch_title.paragraph_format.space_before = Pt(28)
    p_arch_title.paragraph_format.space_after = Pt(6)
    r_at = p_arch_title.add_run("Estructura Modular del Sistema")
    r_at.font.name = "Calibri"
    r_at.font.bold = True
    r_at.font.size = Pt(13)
    r_at.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    p_arch_desc = doc.add_paragraph()
    p_arch_desc.paragraph_format.space_before = Pt(0)
    p_arch_desc.paragraph_format.space_after = Pt(8)
    r_ad = p_arch_desc.add_run(
        "El presente documento consolida la totalidad del código fuente desarrollado para el ecosistema "
        "de gemelos digitales agroecológicos. La solución se organiza bajo una arquitectura desacoplada y orientada "
        "a microservicios que integra cinco pilares fundamentales:\n\n"
        "• Módulo 1 (Infraestructura): Definiciones Docker Compose, variables de entorno y orquestación unificada.\n"
        "• Módulo 2 (Backend API): API REST de alto rendimiento en FastAPI con PostgreSQL, JWT y servicios de optimización.\n"
        "• Módulo 3 (Frontend Web): Aplicación reactiva en React 18, Vite, Tailwind CSS, Leaflet y dioramas 3D en Three.js.\n"
        "• Módulo 4 (Motor ML & Streamlit): Gemelo digital interactivo, modelos biofísicos ABM de polinizadores y tests robustos.\n"
        "• Módulo 5 (Agentes Langflow): Orquestación de agentes de IA con Langflow y pruebas estadísticas post-hoc."
    )
    r_ad.font.name = "Calibri"
    r_ad.font.size = Pt(9.5)
    r_ad.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    doc.add_page_break()

    # ==========================================
    # 2. ÍNDICE COMPLETO DE ARCHIVOS
    # ==========================================
    p_idx_h = doc.add_paragraph()
    p_idx_h.paragraph_format.space_before = Pt(12)
    p_idx_h.paragraph_format.space_after = Pt(6)
    r_ih = p_idx_h.add_run("ÍNDICE GENERAL Y CATÁLOGO DE ARCHIVOS")
    r_ih.font.name = "Calibri"
    r_ih.font.size = Pt(18)
    r_ih.font.bold = True
    r_ih.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    p_idx_sub = doc.add_paragraph()
    p_idx_sub.paragraph_format.space_before = Pt(0)
    p_idx_sub.paragraph_format.space_after = Pt(14)
    r_is = p_idx_sub.add_run(
        "A continuación se presenta el catálogo completo de los 99 archivos que componen la solución tecnológica, "
        "organizados por módulo, indicando su ruta relativa, lenguaje o formato, líneas de código y propósito técnico."
    )
    r_is.font.name = "Calibri"
    r_is.font.size = Pt(9.5)
    r_is.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # Master Index Table
    idx_tbl = doc.add_table(rows=len(file_records) + 1, cols=6)
    idx_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    idx_tbl.autofit = False

    idx_headers = ["N°", "Módulo", "Ruta del Archivo", "Lenguaje", "Líneas", "Propósito / Rol Técnico"]
    idx_widths = [Inches(0.4), Inches(1.1), Inches(2.3), Inches(0.9), Inches(0.6), Inches(1.8)]

    # Header row
    hdr_row = idx_tbl.rows[0]
    for c_idx, title in enumerate(idx_headers):
        cell = hdr_row.cells[c_idx]
        cell.width = idx_widths[c_idx]
        set_cell_margins(cell, top=100, bottom=100, left=60, right=60)
        set_cell_background(cell, "1E3A8A")
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in (0, 3, 4) else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(title)
        r.font.name = "Calibri"
        r.font.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Data rows
    for r_idx, rec in enumerate(file_records):
        row = idx_tbl.rows[r_idx + 1]
        bg = "FFFFFF" if r_idx % 2 == 0 else "F8FAFC"
        
        # Col 0: Index
        c0 = row.cells[0]
        c0.width = idx_widths[0]
        set_cell_margins(c0, top=60, bottom=60, left=50, right=50)
        set_cell_background(c0, bg)
        p = c0.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(r_idx + 1))
        r.font.name = "Calibri"
        r.font.bold = True
        r.font.size = Pt(8.0)
        r.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
        
        # Col 1: Module
        c1 = row.cells[1]
        c1.width = idx_widths[1]
        set_cell_margins(c1, top=60, bottom=60, left=50, right=50)
        set_cell_background(c1, bg)
        p = c1.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(f"Mod. {rec['module_id']}")
        r.font.name = "Calibri"
        r.font.size = Pt(8.0)
        r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        
        # Col 2: Path
        c2 = row.cells[2]
        c2.width = idx_widths[2]
        set_cell_margins(c2, top=60, bottom=60, left=50, right=50)
        set_cell_background(c2, bg)
        p = c2.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(rec["rel_path"])
        r.font.name = "Consolas"
        r.font.bold = True
        r.font.size = Pt(7.5)
        r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        
        # Col 3: Language
        c3 = row.cells[3]
        c3.width = idx_widths[3]
        set_cell_margins(c3, top=60, bottom=60, left=50, right=50)
        set_cell_background(c3, bg)
        p = c3.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(rec["language"])
        r.font.name = "Calibri"
        r.font.size = Pt(8.0)
        r.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
        
        # Col 4: Lines
        c4 = row.cells[4]
        c4.width = idx_widths[4]
        set_cell_margins(c4, top=60, bottom=60, left=50, right=50)
        set_cell_background(c4, bg)
        p = c4.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run(f"{rec['lines_count']:,}")
        r.font.name = "Consolas"
        r.font.size = Pt(8.0)
        r.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
        
        # Col 5: Description
        c5 = row.cells[5]
        c5.width = idx_widths[5]
        set_cell_margins(c5, top=60, bottom=60, left=50, right=50)
        set_cell_background(c5, bg)
        p = c5.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(rec["description"])
        r.font.name = "Calibri"
        r.font.size = Pt(7.5)
        r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    set_table_borders(idx_tbl, "CBD5E1")
    doc.add_page_break()

    # ==========================================
    # 3. LISTADO COMPLETO DE CÓDIGO POR ARCHIVO
    # ==========================================
    global_file_counter = 0
    
    for mod in MODULES:
        # Module Banner Heading
        p_mod = doc.add_paragraph()
        p_mod.paragraph_format.space_before = Pt(24)
        p_mod.paragraph_format.space_after = Pt(6)
        r_mh = p_mod.add_run(mod["name"])
        r_mh.font.name = "Calibri"
        r_mh.font.size = Pt(16)
        r_mh.font.bold = True
        r_mh.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        
        p_mdesc = doc.add_paragraph()
        p_mdesc.paragraph_format.space_before = Pt(0)
        p_mdesc.paragraph_format.space_after = Pt(16)
        r_md = p_mdesc.add_run(mod["description"])
        r_md.font.name = "Calibri"
        r_md.font.size = Pt(9.5)
        r_md.font.italic = True
        r_md.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
        
        # Files in this module
        mod_records = [r for r in file_records if r["module_id"] == mod["id"]]
        
        for rec in mod_records:
            global_file_counter += 1
            print(f"[{global_file_counter}/{len(file_records)}] Agregando: {rec['rel_path']} ({rec['lines_count']} líneas)...")
            
            # File Title Header
            p_fh = doc.add_paragraph()
            p_fh.paragraph_format.space_before = Pt(18)
            p_fh.paragraph_format.space_after = Pt(4)
            p_fh.paragraph_format.keep_with_next = True
            
            r_num = p_fh.add_run(f"Archivo #{global_file_counter:02d} — ")
            r_num.font.name = "Calibri"
            r_num.font.bold = True
            r_num.font.size = Pt(11)
            r_num.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
            
            r_fname = p_fh.add_run(rec["rel_path"])
            r_fname.font.name = "Consolas"
            r_fname.font.bold = True
            r_fname.font.size = Pt(11)
            r_fname.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            
            # File Metadata Box
            meta_tbl = doc.add_table(rows=2, cols=4)
            meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            meta_tbl.autofit = False
            
            m_widths = [Inches(1.7), Inches(1.8), Inches(1.7), Inches(1.9)]
            
            # Row 0
            row0 = meta_tbl.rows[0]
            # Módulo
            c = row0.cells[0]
            c.width = m_widths[0]
            set_cell_margins(c, top=40, bottom=40, left=60, right=60)
            set_cell_background(c, "F1F5F9")
            p = c.paragraphs[0]; p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
            r = p.add_run("Módulo: "); r.font.bold = True; r.font.size = Pt(8); r.font.name = "Calibri"
            r2 = p.add_run(f"Módulo {rec['module_id']}"); r2.font.size = Pt(8); r2.font.name = "Calibri"
            
            # Lenguaje
            c = row0.cells[1]
            c.width = m_widths[1]
            set_cell_margins(c, top=40, bottom=40, left=60, right=60)
            set_cell_background(c, "F1F5F9")
            p = c.paragraphs[0]; p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
            r = p.add_run("Lenguaje: "); r.font.bold = True; r.font.size = Pt(8); r.font.name = "Calibri"
            r2 = p.add_run(rec["language"]); r2.font.size = Pt(8); r2.font.name = "Calibri"
            
            # Líneas
            c = row0.cells[2]
            c.width = m_widths[2]
            set_cell_margins(c, top=40, bottom=40, left=60, right=60)
            set_cell_background(c, "F1F5F9")
            p = c.paragraphs[0]; p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
            r = p.add_run("Líneas: "); r.font.bold = True; r.font.size = Pt(8); r.font.name = "Calibri"
            r2 = p.add_run(f"{rec['lines_count']:,}"); r2.font.size = Pt(8); r2.font.name = "Calibri"
            
            # Tamaño
            c = row0.cells[3]
            c.width = m_widths[3]
            set_cell_margins(c, top=40, bottom=40, left=60, right=60)
            set_cell_background(c, "F1F5F9")
            p = c.paragraphs[0]; p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
            r = p.add_run("Tamaño: "); r.font.bold = True; r.font.size = Pt(8); r.font.name = "Calibri"
            size_kb = rec['size_bytes'] / 1024.0
            r2 = p.add_run(f"{size_kb:.1f} KB ({rec['size_bytes']:,} bytes)"); r2.font.size = Pt(8); r2.font.name = "Calibri"

            # Row 1: Merged Description
            row1 = meta_tbl.rows[1]
            cell_desc = row1.cells[0]
            cell_desc.merge(row1.cells[1]).merge(row1.cells[2]).merge(row1.cells[3])
            set_cell_margins(cell_desc, top=40, bottom=40, left=60, right=60)
            set_cell_background(cell_desc, "F8FAFC")
            p = cell_desc.paragraphs[0]; p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
            r = p.add_run("Propósito: "); r.font.bold = True; r.font.size = Pt(8); r.font.name = "Calibri"
            r2 = p.add_run(rec["description"]); r2.font.size = Pt(8); r2.font.name = "Calibri"

            set_table_borders(meta_tbl, "CBD5E1")
            
            # Small spacer before code block
            p_sp = doc.add_paragraph()
            p_sp.paragraph_format.space_before = Pt(2)
            p_sp.paragraph_format.space_after = Pt(2)
            p_sp.paragraph_format.keep_with_next = True
            
            # Render Code Block
            # We chunk the code into blocks of 60 lines to provide optimal Word pagination
            lines = rec["lines"]
            if not lines:
                lines = ["# (Archivo vacío)"]
            
            chunk_size = 60
            chunks = [lines[i:i + chunk_size] for i in range(0, len(lines), chunk_size)]
            
            for chunk_idx, chunk in enumerate(chunks):
                p_code = doc.add_paragraph()
                pPr = p_code._p.get_or_add_pPr()
                
                # Shading: Soft off-white / light slate #F8FAFC
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC"/>')
                pPr.append(shd)
                
                # Left border: Accent blue #2563EB, 2.25pt
                bdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="18" w:space="8" w:color="2563EB"/></w:pBdr>')
                pPr.append(bdr)
                
                p_code.paragraph_format.space_before = Pt(0)
                p_code.paragraph_format.space_after = Pt(0)
                p_code.paragraph_format.line_spacing = 1.0
                
                # Add text run
                chunk_text = "\n".join(chunk)
                r_c = p_code.add_run(chunk_text)
                r_c.font.name = "Consolas"
                r_c.font.size = Pt(8.0)
                r_c.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            
            # End of file divider / spacing
            p_end = doc.add_paragraph()
            p_end.paragraph_format.space_before = Pt(6)
            p_end.paragraph_format.space_after = Pt(12)

    # Save final Word document
    output_filename = "CODIGO_COMPLETO_PROYECTO_GEMELOS_DIGITALES.docx"
    print(f"Guardando documento final en '{output_filename}'...")
    t_save_0 = time.time()
    doc.save(output_filename)
    t_save_1 = time.time()
    
    # Also save a copy with a simple name
    copy_filename = "CODIGO_COMPLETO_PROYECTO.docx"
    import shutil
    shutil.copyfile(output_filename, copy_filename)
    
    total_time = time.time() - t0
    final_size = Path(output_filename).stat().st_size
    print(f"¡Éxito rotundo!")
    print(f"Documento generado: {output_filename}")
    print(f"Copia alternativa: {copy_filename}")
    print(f"Tamaño final: {final_size / 1024:.1f} KB ({final_size / (1024*1024):.2f} MB)")
    print(f"Tiempo total: {total_time:.1f} segundos (guardado en {t_save_1 - t_save_0:.2f}s)")
    return output_filename

if __name__ == "__main__":
    build_word_document()
