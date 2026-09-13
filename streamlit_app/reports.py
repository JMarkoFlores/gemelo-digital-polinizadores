import io
import pandas as pd
from datetime import datetime

try:
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors
except ImportError:
    pass

try:
    from docx import Document
except ImportError:
    pass


def generate_excel_report(registro_df: pd.DataFrame, comparativa_df: pd.DataFrame, dataset_info_df: pd.DataFrame, raw_dataset: pd.DataFrame, best_model: str, interpretations: dict = None) -> bytes:
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        info_df = pd.DataFrame([
            {"Propiedad": "Fecha de reporte", "Valor": datetime.now().strftime("%Y-%m-%d %H:%M:%S")},
            {"Propiedad": "Mejor Modelo Seleccionado", "Valor": best_model},
        ])
        info_df.to_excel(writer, sheet_name='Info_General', index=False)
        dataset_info_df.to_excel(writer, sheet_name='Resumen_Dataset', index=False)
        raw_dataset.to_excel(writer, sheet_name='Dataset_Completo', index=False)
        registro_df.to_excel(writer, sheet_name='Registro_Modelos', index=False)
        comparativa_df.to_excel(writer, sheet_name='Tabla_Comparativa', index=False)
        
        if interpretations is not None:
            texts = []
            if interpretations.get("b1"): 
                texts.append({"Sección": "T-Student / Wilcoxon", "Interpretación": interpretations["b1"]})
                if interpretations.get("nota_metodologica_b1"):
                    texts.append({"Sección": "T-Student / Wilcoxon (Nota)", "Interpretación": interpretations["nota_metodologica_b1"]})
            if interpretations.get("friedman"): texts.append({"Sección": "Test de Friedman", "Interpretación": interpretations["friedman"]})
            if interpretations.get("nemenyi_text"): texts.append({"Sección": "Test de Nemenyi", "Interpretación": interpretations["nemenyi_text"]})
            if interpretations.get("veredicto"): texts.append({"Sección": "Veredicto Final", "Interpretación": interpretations["veredicto"]})
            
            if texts:
                pd.DataFrame(texts).to_excel(writer, sheet_name='Interpretaciones', index=False)
                
            if interpretations.get("nemenyi_matrix") is not None:
                interpretations["nemenyi_matrix"].to_excel(writer, sheet_name='Nemenyi_Matrix')
            
    return output.getvalue()


def _add_df_to_word(doc, df, title, include_index=False):
    doc.add_heading(title, level=1)
    df_stringified = df.copy()
    
    if include_index:
        df_stringified.insert(0, 'Modelo', df_stringified.index)
        
    for col in df_stringified.columns:
        if df_stringified[col].dtype.kind in 'fc':
            df_stringified[col] = df_stringified[col].round(4)
            
    table = doc.add_table(rows=1, cols=len(df_stringified.columns))
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    for i, col in enumerate(df_stringified.columns):
        hdr_cells[i].text = str(col)
        
    for index, row in df_stringified.iterrows():
        row_cells = table.add_row().cells
        for i, val in enumerate(row):
            row_cells[i].text = str(val)
    doc.add_paragraph("")


def generate_word_report(registro_df: pd.DataFrame, comparativa_df: pd.DataFrame, dataset_info_df: pd.DataFrame, best_model: str, interpretations: dict = None) -> bytes:
    doc = Document()
    doc.add_heading('Reporte de Evaluación de Modelos IA', 0)
    
    doc.add_paragraph(f'Fecha de generación: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}')
    p = doc.add_paragraph('El mejor modelo seleccionado fue: ')
    p.add_run(best_model).bold = True
    
    _add_df_to_word(doc, dataset_info_df, 'Cantidad de Datos y Campos')
    _add_df_to_word(doc, registro_df, 'Registro de Modelos')
    _add_df_to_word(doc, comparativa_df, 'Tabla Comparativa de Modelos')
    
    if interpretations is not None:
        doc.add_heading('Interpretaciones Estadísticas', level=1)
        if interpretations.get("b1"):
            doc.add_paragraph("T-Student / Wilcoxon: " + interpretations["b1"])
            if interpretations.get("nota_metodologica_b1"):
                p_nota = doc.add_paragraph(interpretations["nota_metodologica_b1"])
                p_nota.style = 'Intense Quote'
        if interpretations.get("friedman"):
            doc.add_paragraph("Test de Friedman: " + interpretations["friedman"])
        if interpretations.get("nemenyi_text"):
            doc.add_paragraph("Test de Nemenyi: " + interpretations["nemenyi_text"])
        if interpretations.get("veredicto"):
            doc.add_paragraph("Veredicto Final:\n" + interpretations["veredicto"].replace("**", ""))
            
        if interpretations.get("nemenyi_matrix") is not None:
            _add_df_to_word(doc, interpretations["nemenyi_matrix"], 'Matriz de p-valores (Nemenyi)', include_index=True)
    
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def _add_df_to_pdf(elements, df, title, styles, t_style, include_index=False):
    elements.append(Paragraph(title, styles['Heading2']))
    
    df_to_plot = df.copy()
    if include_index:
        df_to_plot.insert(0, 'Modelo', df_to_plot.index)
        
    data_list = [df_to_plot.columns.tolist()]
    for index, row in df_to_plot.iterrows():
        formatted_row = []
        for val in row:
            if isinstance(val, float):
                formatted_row.append(f"{val:.4f}")
            else:
                formatted_row.append(str(val))
        data_list.append(formatted_row)
        
    table = Table(data_list)
    table.setStyle(t_style)
    elements.append(table)
    elements.append(Spacer(1, 20))


def generate_pdf_report(registro_df: pd.DataFrame, comparativa_df: pd.DataFrame, dataset_info_df: pd.DataFrame, best_model: str, interpretations: dict = None) -> bytes:
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(letter))
    elements = []
    
    styles = getSampleStyleSheet()
    title_style = styles['Title']
    normal_style = styles['Normal']
    
    elements.append(Paragraph("Reporte de Evaluación de Modelos IA", title_style))
    elements.append(Paragraph(f'Fecha: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}', normal_style))
    elements.append(Paragraph(f'<b>Mejor modelo seleccionado: {best_model}</b>', normal_style))
    elements.append(Spacer(1, 20))
    
    t_style = TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#2a3f5f")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 8),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
        ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor("#f4f4f4")),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 7),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey)
    ])
    
    _add_df_to_pdf(elements, dataset_info_df, "Cantidad de Datos y Campos Entrenados", styles, t_style)
    _add_df_to_pdf(elements, registro_df, "Registro de Modelos", styles, t_style)
    _add_df_to_pdf(elements, comparativa_df, "Tabla Comparativa de Modelos", styles, t_style)
    
    if interpretations is not None:
        elements.append(Paragraph("Interpretaciones Estadísticas", styles['Heading2']))
        if interpretations.get("b1"):
            elements.append(Paragraph("<b>T-Student / Wilcoxon:</b> " + interpretations["b1"].replace("**", ""), normal_style))
            if interpretations.get("nota_metodologica_b1"):
                elements.append(Paragraph("<i>" + interpretations["nota_metodologica_b1"] + "</i>", normal_style))
            elements.append(Spacer(1, 10))
        if interpretations.get("friedman"):
            elements.append(Paragraph("<b>Test de Friedman:</b> " + interpretations["friedman"].replace("**", ""), normal_style))
            elements.append(Spacer(1, 10))
        if interpretations.get("nemenyi_text"):
            elements.append(Paragraph("<b>Test de Nemenyi:</b> " + interpretations["nemenyi_text"].replace("**", ""), normal_style))
            elements.append(Spacer(1, 10))
        if interpretations.get("veredicto"):
            # Replace double newlines with single spaces or break tags if needed, but Paragraph handles some formatting.
            veredicto_clean = interpretations["veredicto"].replace("\n\n", "<br/>").replace("**", "")
            elements.append(Paragraph("<b>Veredicto Final:</b><br/>" + veredicto_clean, normal_style))
            elements.append(Spacer(1, 10))
            
        if interpretations.get("nemenyi_matrix") is not None:
            _add_df_to_pdf(elements, interpretations["nemenyi_matrix"], "Matriz de p-valores (Nemenyi)", styles, t_style, include_index=True)
            
    doc.build(elements)
    
    return output.getvalue()
